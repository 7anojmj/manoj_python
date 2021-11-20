# Line Chart
import random
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference

wb = Workbook()
spreadsheet = wb.active

# Let's create some sample data
rows = [
    ["", "January", "February", "March", "April",
     "May", "June", "July", "August", "September",
     "October", "November", "December"],
    [1, ],
    [2, ],
    [3, ],
]

for row in rows:
    spreadsheet.append(row)

for row in spreadsheet.iter_rows(min_row=2,
                                 max_row=4,
                                 min_col=2,
                                 max_col=13):
    for cell in row:
        cell.value = random.randrange(5, 100)

chart = LineChart()
data = Reference(worksheet=spreadsheet,
                 min_row=2,
                 max_row=4,
                 min_col=1,
                 max_col=13)

chart.add_data(data, from_rows=True, titles_from_data=True)
spreadsheet.add_chart(chart, "C6")

wb.save("line_chart1.xlsx")
